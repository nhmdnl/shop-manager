"""Round-trips app strings through a spreadsheet for a translator.

    uv run --with openpyxl tools/translation_sheet.py export ti translations/Arch-SME_Tigrinya.xlsx
    uv run --with openpyxl tools/translation_sheet.py import ti translations/Arch-SME_Tigrinya.xlsx

export: one row per English string (plurals one row per quantity), with where it appears and the current
translation pre-filled. import: rewrites values-<lang>/strings.xml from the sheet in English order; a blank
cell is left out, so that string falls back to English.
"""

import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
RES = ROOT / "app/src/main/res"
SRC = ROOT / "app/src/main/java/com/archsolutions/shopmanager"
LANGUAGES = {"ti": "Tigrinya", "lg": "Luganda", "sw": "Kiswahili", "pt": "Portuguese"}
FEATURES = {
    "home": "Home", "sell": "Sell", "sales": "Today's sales", "stock": "Stock", "suppliers": "Suppliers",
    "day": "Day", "reports": "Reports", "settings": "Settings", "pin": "PIN", "onboarding": "Onboarding",
    "designsystem": "Shared", "app": "App", "domain": "Messages", "data": "Backup",
}
# File stems named the way the screen reads.
PLACES = {
    "AppNavHost": "Tab bar", "Backup": "Backup and restore", "Checkout": "Checkout", "CloseReminder": "Close-day reminder (notification)", "CloseDay": "Close the day",
    "DayClosed": "Day closed", "DayGateLockCard": "Locked day card", "Debtors": "Owed to the shop",
    "NoteCounter": "Note counter", "OpenDay": "Open the day", "PastDay": "Past day", "RecordMoney": "Record money",
    "RecordMoneyForm": "Record money", "ReportExport": "Export", "SaleParts": "Cart", "ProductImport": "Product import",
    "DayAndStats": "Day bar", "States": "Empty and error states", "SteppersAndTiles": "Product tiles",
    "AddPurchase": "Add purchase", "PurchaseForm": "Add purchase", "EditProduct": "Edit product",
    "SupplierDetail": "Supplier detail", "SupplierEdit": "Edit supplier", "SaleDetail": "Sale detail", "Sales": "Today's sales",
}
HEADERS = ["#", "Section", "Where it appears", "English", "Translation", "Notes", "Key"]
PLACEHOLDER = re.compile(r"%(\d+\$)?[sd]|%%")


def unescape(text):
    """Android resource text as the person sees it."""
    text = text.strip()
    if len(text) >= 2 and text[0] == text[-1] == '"':
        text = text[1:-1]
    return re.sub(r"\\(.)", lambda m: "\n" if m.group(1) == "n" else m.group(1), text)


def escape(text):
    text = text.replace("\\", "\\\\").replace("\r\n", "\n").replace("\n", "\\n")
    text = text.replace("'", "\\'").replace('"', '\\"').replace("&", "&amp;").replace("<", "&lt;")
    return "\\" + text if text[:1] in ("@", "?") else text


def read_strings(path):
    """[(section, key, quantity or None, text)] in file order; untranslatable strings left out."""
    entries, section = [], ""
    if not path.exists():
        return entries
    # Our own resource files, not outside input.
    tree = ET.parse(path, parser=ET.XMLParser(target=ET.TreeBuilder(insert_comments=True)))
    for node in tree.getroot():
        if node.tag is ET.Comment:
            if "DRAFT" not in (node.text or ""):
                section = node.text.strip()
        elif node.tag == "string" and node.get("translatable") != "false":
            entries.append((section, node.get("name"), None, unescape(node.text or "")))
        elif node.tag == "plurals":
            for item in node:
                entries.append((section, node.get("name"), item.get("quantity"), unescape(item.text or "")))
    return entries


def usages():
    """key -> sorted screen names that reference R.string.key / R.plurals.key."""
    found = {}
    for file in SRC.rglob("*.kt"):
        parts = file.relative_to(SRC).parts
        area = next((FEATURES[p] for p in parts if p in FEATURES), "App")
        stem = re.sub(r"(Screen|ViewModel|Panel|Sheet|Dialog)$", "", file.stem)
        place = PLACES.get(stem, area if stem.lower() == area.lower().replace(" ", "").replace("'", "") else f"{area} › {stem}")
        for key in re.findall(r"R\.(?:string|plurals)\.(\w+)", file.read_text(encoding="utf-8")):
            found.setdefault(key, set()).add(place)
    def short(places):
        places = sorted(places)
        return ", ".join(places[:3]) + (f" and {len(places) - 3} more" if len(places) > 3 else "")
    return {key: short(places) for key, places in found.items()}


def notes(key, quantity, english):
    out = []
    marks = sorted(set(m.group(0) for m in PLACEHOLDER.finditer(english)))
    if marks:
        out.append("Keep " + " ".join(marks) + " (filled in by the app; you can move it)")
    if "%%" in marks:
        out.append("%% shows as %")
    if key.startswith("nav_"):
        out.append("Tab label: keep it short")
    elif key.startswith("action_"):
        out.append("Button label")
    if quantity:
        out.append(f"Plural: '{quantity}' form" + (" (used for 0 and 1 in Tigrinya)" if quantity == "one" else ""))
    if "\n" in english:
        out.append("Keep the line break")
    return "; ".join(out)


def export(lang, out):
    name = LANGUAGES[lang]
    english = read_strings(RES / "values/strings.xml")
    current = {(k, q): t for _, k, q, t in read_strings(RES / f"values-{lang}/strings.xml")}
    where = usages()

    wb = Workbook()
    guide = wb.active
    guide.title = "How to fill"
    lines = [
        (f"Arch-SME: {name} translation", Font(bold=True, size=14)),
        ("", None),
        (f"1. Open the 'Strings' tab. Write your {name} in the yellow '{name}' column. Only that column is read back.", None),
        (f"2. {sum(1 for k in current.values() if k)} rows already hold a draft. Correct it, or leave it if it is right.", None),
        ("3. A blank cell shows English in the app, so you can stop half way and send what you have.", None),
        ("4. Keep %1$s, %2$d and similar marks exactly as written. The app puts an amount, a name or a number there; move the mark to wherever it belongs in your sentence.", None),
        ("5. Amounts and numbers stay in Latin digits (UGX 5,000). Never write a number yourself where a % mark is.", None),
        ("6. Tab labels and buttons must be short: a phone shows five tabs side by side.", None),
        ("7. 'Where it appears' names the screen, so you can open it in the app and see the words in place.", None),
        ("8. Rows marked [one] and [other] are the singular and plural forms of the same text. Fill both.", None),
        ("9. Save the file (keep it .xlsx) and send it back. It is imported into the app as it is.", None),
    ]
    for row, (text, font) in enumerate(lines, start=1):
        cell = guide.cell(row=row, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            cell.font = font
    guide.column_dimensions["A"].width = 110

    sheet = wb.create_sheet("Strings")
    sheet.append(HEADERS[:4] + [name] + HEADERS[5:])
    for n, (section, key, quantity, text) in enumerate(english, start=1):
        label = f"{key} [{quantity}]" if quantity else key
        sheet.append([n, section, where.get(key, "Not shown in the app yet: skip"), text, current.get((key, quantity), ""), notes(key, quantity, text), label])

    widths = {"A": 6, "B": 22, "C": 30, "D": 48, "E": 48, "F": 36, "G": 30}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    yellow, grey = PatternFill("solid", fgColor="FFF4C2"), Font(color="808080")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = yellow
        row[4].font = Font(name="Noto Sans Ethiopic", size=12) if lang == "ti" else Font(size=12)
        row[3].font = Font(size=12)
        row[6].font = grey
    table = Table(displayName="Strings", ref=f"A1:G{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = "E2"
    # Printable too: landscape, one page wide, the header on every page.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    wb.active = 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{out}: {len(english)} rows, {sum(1 for v in current.values() if v)} pre-filled")


def import_sheet(lang, path):
    sheet = load_workbook(path, read_only=True)["Strings"]
    header = [c.value for c in next(sheet.iter_rows(max_row=1))]
    key_col, text_col = header.index("Key"), 4
    done = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        label, text = row[key_col], row[text_col]
        if not label or text is None or not str(text).strip():
            continue
        m = re.fullmatch(r"(\w+)(?: \[(\w+)\])?", label)
        done[(m.group(1), m.group(2))] = str(text).strip()

    english = read_strings(RES / "values/strings.xml")
    problems = []
    for _, key, quantity, text in english:
        if (key, quantity) in done:
            want = sorted(m.group(0) for m in PLACEHOLDER.finditer(text))
            got = sorted(m.group(0) for m in PLACEHOLDER.finditer(done[(key, quantity)]))
            if want != got:
                problems.append(f"  {key}{f' [{quantity}]' if quantity else ''}: needs {' '.join(want) or 'no marks'}, has {' '.join(got) or 'none'}")
    if problems:
        sys.exit("Placeholders differ from English, nothing written:\n" + "\n".join(problems))

    out, section, written = ['<?xml version="1.0" encoding="utf-8"?>', "<resources>"], None, 0
    plurals_open = None
    for sec, key, quantity, _ in english:
        if (key, quantity) not in done:
            continue
        if plurals_open and plurals_open != key:
            out.append("    </plurals>")
            plurals_open = None
        if sec != section:
            out += ["", f"    <!-- {sec} -->"]
            section = sec
        value = escape(done[(key, quantity)])
        if quantity:
            if plurals_open != key:
                out.append(f'    <plurals name="{key}">')
                plurals_open = key
            out.append(f'        <item quantity="{quantity}">{value}</item>')
        else:
            out.append(f'    <string name="{key}">{value}</string>')
        written += 1
    if plurals_open:
        out.append("    </plurals>")
    out.append("</resources>")
    target = RES / f"values-{lang}/strings.xml"
    target.parent.mkdir(exist_ok=True)
    target.write_text("\n".join(out) + "\n", encoding="utf-8")
    print(f"{target}: {written} of {len(english)} strings")


if __name__ == "__main__":
    if len(sys.argv) != 4 or sys.argv[1] not in ("export", "import") or sys.argv[2] not in LANGUAGES:
        sys.exit(__doc__)
    (export if sys.argv[1] == "export" else import_sheet)(sys.argv[2], Path(sys.argv[3]).expanduser())
